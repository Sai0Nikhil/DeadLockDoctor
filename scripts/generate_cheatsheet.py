#!/usr/bin/env python3
"""
Generates DeadlockDoctor_Commands_CheatSheet.docx
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_commands_docx():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    NAVY = RGBColor(16, 44, 87)
    CYAN = RGBColor(0, 120, 160)
    DARK = RGBColor(30, 35, 45)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("DEADLOCKDOCTOR: COMMANDS & SHELL MANUAL")
    r.font.name = 'Calibri'; r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Complete Command Reference, Terminal Flags & Interactive REPL Cheat Sheet")
    r.font.name = 'Calibri'; r.font.size = Pt(12); r.font.color.rgb = CYAN

    commands = [
        ("Binary Flag", "./deadlockdoctor --shell (-s)", "Launch directly into the interactive DeadlockDoctor Shell (REPL)."),
        ("Binary Flag", "./deadlockdoctor --demo (-d)", "Run the real POSIX Multithreading Mutex Race & Watchdog demo."),
        ("Binary Flag", "./deadlockdoctor --walkthrough (-w)", "Run the educational step-by-step avoidance and recovery walkthrough."),
        ("Shell Command", "show / matrices", "Display Allocation, Need, Max, and Available matrices in ANSI color tables."),
        ("Shell Command", "graph / graphs", "Display ASCII Resource Allocation Graph (RAG) and Wait-For Graph (WFG)."),
        ("Shell Command", "banker", "Run Dijkstra's Banker's Algorithm to check safety and compute safe sequence."),
        ("Shell Command", "detect", "Run 3-Color DFS cycle detection on WFG and multi-instance reduction."),
        ("Shell Command", "recover", "Execute starvation-aware cost-based victim selection and unblock dependents."),
        ("Shell Command", "add_res <name> <inst>", "Register a new resource type with total instances (e.g. add_res GPU 4)."),
        ("Shell Command", "add_proc <name> [prio]", "Register a new process with priority (e.g. add_proc Worker1 5)."),
        ("Shell Command", "set_max <p> <r> <k>", "Set maximum claim of process p on resource r to k units."),
        ("Shell Command", "req <p> <r> <k>", "Process p requests k units of resource r (evaluated via mode)."),
        ("Shell Command", "rel <p> <r> <k>", "Process p releases k units of resource r back to available pool."),
        ("Shell Command", "mode <avoid|detect|prev>", "Switch operating mode between Avoidance, Detection, and Prevention."),
        ("OS Command", "pwd, ls, clear, cd, whoami", "Executed seamlessly directly via native host operating system subshell."),
        ("Build Target", "make all / make test / make memcheck", "Build binaries, run 24-assertion test suite, and run AddressSanitizer check.")
    ]

    tbl = doc.add_table(rows=1, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = tbl.rows[0]
    hdr.cells[0].width = Inches(1.5); hdr.cells[1].width = Inches(2.3); hdr.cells[2].width = Inches(3.1)
    for idx, name in enumerate(["Category", "Command Syntax", "Description & Action"]):
        set_cell_background(hdr.cells[idx], "102C57")
        p = hdr.cells[idx].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        r = p.add_run(name); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255); r.font.size = Pt(10)

    for cat, cmd, desc in commands:
        row = tbl.add_row()
        row.cells[0].width = Inches(1.5); row.cells[1].width = Inches(2.3); row.cells[2].width = Inches(3.1)
        set_cell_background(row.cells[0], "F8FAFC"); set_cell_background(row.cells[1], "F0F4F8"); set_cell_background(row.cells[2], "FFFFFF")
        for i, val in enumerate([cat, cmd, desc]):
            p = row.cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val); r.font.size = Pt(9.5)
            if i == 0: r.font.bold = True; r.font.color.rgb = CYAN
            if i == 1: r.font.bold = True; r.font.color.rgb = DARK

    out = "docs/reports/DeadlockDoctor_Commands_CheatSheet.docx"
    doc.save(out)
    print(f"[+] Word Cheat Sheet generated: {out}")

if __name__ == "__main__":
    generate_commands_docx()
